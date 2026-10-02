# knowledge_base/document_loader.py

from pathlib import Path
from typing import List, Dict

from pypdf import PdfReader
from pptx import Presentation
from docx import Document


def load_pdf(file_path: str) -> str:
    """读取 PDF 文本"""
    reader = PdfReader(file_path)
    texts = []

    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)

    return "\n".join(texts)


def load_pptx(file_path: str) -> str:
    """读取 PPT 文本"""
    prs = Presentation(file_path)
    texts = []

    for i, slide in enumerate(prs.slides, start=1):
        slide_text = [f"第{i}页："]

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

        texts.append("\n".join(slide_text))

    return "\n".join(texts)


def load_docx(file_path: str) -> str:
    """读取 Word 文本"""
    doc = Document(file_path)
    texts = []

    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text.strip())

    return "\n".join(texts)


def load_txt(file_path: str) -> str:
    """读取 TXT 文本"""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def load_document(file_path: str) -> Dict:
    """
    根据文件类型自动读取文档
    返回格式：
    {
        "file_name": 文件名,
        "file_path": 文件路径,
        "content": 文本内容
    }
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        content = load_pdf(file_path)
    elif suffix == ".pptx":
        content = load_pptx(file_path)
    elif suffix == ".docx":
        content = load_docx(file_path)
    elif suffix == ".txt":
        content = load_txt(file_path)
    else:
        raise ValueError(f"暂不支持该文件类型：{suffix}")

    return {
        "file_name": path.name,
        "file_path": str(path),
        "content": content
    }


def load_documents_from_folder(folder_path: str) -> List[Dict]:
    """
    读取一个文件夹下的所有课程资料
    支持：pdf、pptx、docx、txt
    """
    folder = Path(folder_path)

    supported_suffixes = [".pdf", ".pptx", ".docx", ".txt"]
    documents = []

    for file in folder.iterdir():
        if file.suffix.lower() in supported_suffixes:
            try:
                doc = load_document(str(file))
                documents.append(doc)
                print(f"[OK] 已读取：{file.name}")
            except Exception as e:
                print(f"[ERROR] 读取失败：{file.name}，原因：{e}")

    return documents